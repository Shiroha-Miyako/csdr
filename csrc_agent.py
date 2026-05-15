import argparse
import json
import os
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from docx import Document
from openpyxl import load_workbook
# 某些证监会表格或手工模板里带有 Excel 图表/绘图 XML，openpyxl 读取时可能报 pitchFamily 错误。
# 这个工具只处理单元格数据，不需要读取图片/图表，因此禁用图片/图表解析。
try:
    import openpyxl.reader.excel as _openpyxl_excel
    _openpyxl_excel.find_images = lambda archive, path: ([], [])
except Exception:
    pass
from rapidfuzz import fuzz, process

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass

CHINESE_NUM = "一二三四五六七八九十"

# 你可以持续维护这个词典：优先级高于 AI
KEYWORD_RULES = [
    (["股权架构", "返程并购", "37 号文", "37号文"], "股权架构搭建及返程并购的合规性", "业务合规"),
    (["业务经营", "数字广告", "网络推广", "职业中介", "广播电视节目制作"], "业务经营合规性", "业务合规"),
    (["境内运营实体", "重大诉讼", "行政处罚", "历次股权变动"], "境内运营实体情况", "业务合规"),
    (["高耗能", "高排放", "排污", "固定污染物"], "已建、在建及此次募投项目是否属于“高耗能”“高排放”项目", "环保"),
    (["全流通", "质押", "冻结", "权利瑕疵"], "全流通股份的瑕疵情形", "全流通"),
    (["国有股东", "国有股标识", "国资"], "国有股东标识办理进展情况", "国有股份"),
    (["股权激励", "员工持股", "外部顾问", "期权激励"], "股权激励计划", "股权激励"),
    (["最近12个月", "新增股东", "入股价格", "利益输送"], "新增股东入股价格的定价依据及合理性", "股东情况"),
    (["股份代持", "股权代持", "代持"], "股份代持情况", "股份代持"),
    (["历次增资", "股权转让", "实缴出资", "抽逃出资", "出资方式"], "历次股权变动的合法合规性", "股权变更"),
    (["持股5%以上", "5%以上", "穿透", "禁止持股"], "持股5%以上股东信息", "股东情况"),
    (["董事", "高级管理人员", "国籍", "永久境外居留权"], "董事、高级管理人员国籍及境外永久居留权情况", "股东情况"),
    (["APP", "小程序", "公众号", "个人信息", "数据收集", "数据安全"], "主要境内运营实体开发、运营的网站、APP等产品情况；个人信息保护、数据安全情况", "数据安全"),
    (["境外子公司", "境外投资", "外汇登记"], "公司境外子公司涉及的境外投资、外汇登记等监管程序具体履行情况说明", "业务合规"),
    (["外商投资准入", "负面清单", "外资禁止", "外资限制"], "公司及下属公司经营范围是否涉及国家禁止或限制外商投资的领域", "业务合规"),
    (["技术进出口", "技术出口"], "最近三年技术出口业务的开展情况及合规性说明", "业务合规"),
    (["募集资金", "募投", "境外投资"], "募集资金用途涉及境外投资所履行的相关审批、核准或备案程序情况", "备案情况"),
    (["发行方案", "招股说明书", "备案材料"], "本次发行上市方案", "上市方案"),
    (["审计意见", "会计师事务所"], "会计师事务所审计意见类型", "备案情况"),
    (["A股上市", "辅导备案", "撤回原因"], "曾经申请A股上市的相关情况以及是否存在对本次发行上市产生重大影响的事项", "历史上市"),
    (["诉讼", "仲裁"], "未决诉讼情况", "未决诉讼"),
    (["安全生产", "行政处罚"], "安全生产法律法规落实情况", "生产安全"),
    (["第八条", "不得境外发行上市", "禁止境外发行上市"], "专项核查（是否存在不得境外发行上市的情形）", "法律合规"),
]



class WorkbookReadError(ValueError):
    """用户上传的 Excel 文件无法被 openpyxl 读取时使用的友好错误。"""


def is_valid_xlsx(path: str) -> bool:
    """xlsx 本质是 zip 包。先用 zipfile 做快速校验，避免云端报红色 ValueError。"""
    try:
        return zipfile.is_zipfile(path)
    except Exception:
        return False


def safe_load_workbook(path: str, **kwargs):
    """
    更稳地读取 Excel 工作簿。
    常见失败原因：文件不是标准 .xlsx、下载成 HTML、.xls 改后缀、文件损坏、外部链接/绘图 XML 兼容问题。
    """
    path = str(path)
    if not is_valid_xlsx(path):
        raise WorkbookReadError(
            "上传的文件不是有效的 .xlsx 工作簿。请用 Excel 或 WPS 打开该文件，"
            "选择“另存为”→“Excel 工作簿 (*.xlsx)”，再重新上传。"
        )

    options = {
        "data_only": False,
        "read_only": False,
        "keep_vba": False,
        "keep_links": False,
    }
    options.update(kwargs)

    try:
        return load_workbook(path, **options)
    except TypeError:
        # 兼容不同 openpyxl 版本的参数差异
        options.pop("rich_text", None)
        try:
            return load_workbook(path, **options)
        except Exception as e:
            raise WorkbookReadError(
                "openpyxl 无法读取该 Excel 文件。请先用 Excel/WPS 打开并另存为新的 .xlsx 文件，"
                "文件名尽量使用英文，例如 master.xlsx、filing_table.xlsx。"
            ) from e
    except Exception as e:
        raise WorkbookReadError(
            "openpyxl 无法读取该 Excel 文件。请先用 Excel/WPS 打开并另存为新的 .xlsx 文件，"
            "文件名尽量使用英文，例如 master.xlsx、filing_table.xlsx。"
        ) from e




def format_date_slash(value) -> str:
    """把日期统一格式化为 2026/5/8，不使用 2026-05-08 或前导 0。"""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value.year}/{value.month}/{value.day}"
    text = str(value).strip()
    if not text:
        return ""
    # 处理 Excel 日期字符串可能带时间：2026-05-08 00:00:00
    text = text.split()[0]
    patterns = [
        r"^(\d{4})[-/\.](\d{1,2})[-/\.](\d{1,2})$",
        r"^(\d{4})年\s*(\d{1,2})月\s*(\d{1,2})日?$",
    ]
    for pat in patterns:
        m = re.search(pat, text)
        if m:
            y, mo, d = m.groups()
            return f"{int(y)}/{int(mo)}/{int(d)}"
    return text.replace("-", "/")


def extract_supplement_dates_from_filename(filename: str) -> Tuple[str, str]:
    """
    从补充材料 Word 文件名提取：
    境外发行上市备案补充材料要求公示（2026年4月27日—2026年5月8日）.docx
    -> (2026/4/27, 2026/5/8)
    第一个日期写入 J列“备案补充材料公告当周”，第二个日期写入 K列“备案补充材料公告日期”。
    """
    name = Path(filename).name
    name = name.replace("（", "(").replace("）", ")")
    name = name.replace("—", "-").replace("–", "-").replace("－", "-").replace("至", "-").replace("到", "-")
    # 支持：2026年4月27日-2026年5月8日 / 2026年4月27日-5月8日
    m = re.search(
        r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日\s*-\s*(?:(\d{4})\s*年\s*)?(\d{1,2})\s*月\s*(\d{1,2})\s*日",
        name,
    )
    if not m:
        return "", ""
    y1, m1, d1, y2, m2, d2 = m.groups()
    y2 = y2 or y1
    return f"{int(y1)}/{int(m1)}/{int(d1)}", f"{int(y2)}/{int(m2)}/{int(d2)}"


@dataclass
class Issue:
    company: str
    original: str
    summary: str = ""
    category: str = ""


def clean_text(s: str) -> str:
    s = (s or "").replace("\u3000", " ").replace("\xa0", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n+", "\n", s)
    return s.strip()


def read_docx_text(path: str) -> str:
    doc = Document(path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n".join(paras)
    # 去掉独立页码行
    text = re.sub(r"\n\d+\n", "\n", text)
    return clean_text(text)


def split_company_sections(text: str) -> Dict[str, str]:
    """识别 docx 中的公司段落。公司名通常是单独一行，后面跟“请你公司...”"""
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    start_idx = 0
    for i, line in enumerate(lines):
        if "具体如下" in line:
            start_idx = i + 1
            break
    lines = lines[start_idx:]

    sections: Dict[str, List[str]] = {}
    current_company = None
    buf: List[str] = []

    def is_company_line(line: str, next_line: str = "") -> bool:
        if len(line) > 40:
            return False
        if re.search(r"[，。；：、（）()]", line):
            return False
        if line.startswith(tuple(CHINESE_NUM)):
            return False
        if "请你公司" in next_line or "请发行人" in next_line:
            return True
        # 部分公告公司名后一行可能不是“请你公司”，保守兜底
        return bool(re.search(r"(电子|网络|通信|科技|新能源|医疗|医药|控股|集团|股份|有限|公司|生物|材料|智能|环保)$", line)) and len(line) <= 25

    for i, line in enumerate(lines):
        nxt = lines[i+1] if i+1 < len(lines) else ""
        if is_company_line(line, nxt):
            if current_company and buf:
                sections[current_company] = buf[:]
            current_company = line
            buf = []
        elif current_company:
            buf.append(line)
    if current_company and buf:
        sections[current_company] = buf

    return {k: clean_text("\n".join(v)) for k, v in sections.items()}


def remove_intro(section_text: str) -> str:
    # 删除“请你公司补充说明...法律意见”等引导句，保留实质问题
    section_text = re.sub(r"^请你公司(?:就以下事项)?补充说明以下事项[，,]?请律师(?:进行)?核查并出具明确的?法律意见[:：]?", "", section_text)
    section_text = re.sub(r"^请你公司就以下事项补充说明[，,]?请律师(?:进行)?核查并出具明确的?法律意见[:：]?", "", section_text)
    section_text = re.sub(r"^请你公司补充说明以下事项[，,]?请律师进行核查并出具明确的法律意见[:：]?", "", section_text)
    return clean_text(section_text)


def split_big_issues(section_text: str) -> List[str]:
    body = remove_intro(section_text)
    # 标记“一、二、三”大序号；不匹配（1）（2）
    pattern = re.compile(rf"(?<![（(])([{CHINESE_NUM}]{{1,3}})、")
    matches = list(pattern.finditer(body))
    if not matches:
        return [clean_text(body)] if body else []
    issues = []
    for idx, m in enumerate(matches):
        start = m.start()
        end = matches[idx+1].start() if idx + 1 < len(matches) else len(body)
        q = body[start:end]
        q = re.sub(rf"^[{CHINESE_NUM}]{{1,3}}、", "", q).strip()
        issues.append(clean_text(q))
    return [x for x in issues if x]


def load_history_examples(wb) -> List[Tuple[str, str, str]]:
    """从问题类型统计、本次新增、当期新增读取历史：原问题/标准问题/大类。"""
    examples = []
    for ws_name in ["问题类型统计", "本次新增", "当期新增"]:
        if ws_name not in wb.sheetnames:
            continue
        ws = wb[ws_name]
        for row in ws.iter_rows(min_row=1, values_only=True):
            vals = list(row)
            if ws_name == "问题类型统计":
                original, category = vals[0], vals[1]
                summary = vals[0]
            else:
                original = vals[1] if len(vals) > 1 else None
                summary = vals[2] if len(vals) > 2 else None
                category = vals[3] if len(vals) > 3 else None
            if original and summary and category and str(original).strip() not in ["问题", "请说明"]:
                examples.append((str(original).strip(), str(summary).strip(), str(category).strip()))
    return examples


def classify_by_rules(issue_text: str, examples: List[Tuple[str, str, str]]) -> Tuple[str, str, str]:
    # 1) 关键词规则
    for kws, summary, category in KEYWORD_RULES:
        issue_no_space = re.sub(r"\s+", "", issue_text)
        if any(k in issue_text or re.sub(r"\s+", "", k) in issue_no_space for k in kws):
            return summary, category, "keyword"
    # 2) 历史相似问题匹配
    if examples:
        candidates = [x[0] for x in examples]
        best = process.extractOne(issue_text, candidates, scorer=fuzz.token_set_ratio)
        if best and best[1] >= 72:
            idx = best[2]
            return examples[idx][1], examples[idx][2], f"history:{best[1]:.0f}"
    # 3) 兜底
    return "待人工确认", "待分类", "fallback"


def parse_supplement_doc(docx_path: str, master_wb_path: str) -> List[Issue]:
    text = read_docx_text(docx_path)
    sections = split_company_sections(text)
    wb = safe_load_workbook(master_wb_path)
    examples = load_history_examples(wb)

    issues: List[Issue] = []
    for company, sec in sections.items():
        for q in split_big_issues(sec):
            summary, cat, _ = classify_by_rules(q, examples)
            issues.append(Issue(company=company, original=q, summary=summary, category=cat))
    return issues


def find_header_row(ws, header="企业名称") -> int:
    for row in range(1, min(ws.max_row, 20) + 1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row, col).value == header:
                return row
    return 3



# ---------------- 公司名多轮差分匹配 ----------------
# 不建议手工维护“简称 -> 全称”表，因此这里用多轮匹配：
# 1) 精确匹配；2) 互相包含；3) 去掉行业/后缀后的核心品牌词匹配；
# 4) 字符覆盖率 + 最长公共子序列兜底。
COMPANY_LEGAL_SUFFIXES = [
    "股份有限公司", "有限责任公司", "有限公司", "集团股份", "集团", "控股", "公司"
]

# 这些词常出现在公告简称末尾或全称中，用于提取“核心品牌词”。
# 例如：迈瑞医疗 -> 迈瑞；铂科电子 -> 铂科；绿联科技 -> 绿联。
COMPANY_INDUSTRY_SUFFIXES = [
    "科技", "技术", "电子", "网络", "通信", "通讯", "新能源", "新材料", "材料",
    "医疗", "医药", "生物", "环保", "节能", "智能", "信息", "股份", "有限", "控股", "集团"
]

COMMON_LOCATION_PREFIXES = [
    "深圳市", "深圳", "上海市", "上海", "北京市", "北京", "广州市", "广州", "南京市", "南京",
    "成都市", "成都", "杭州", "杭州市", "苏州", "苏州市", "香港", "中国"
]


def normalize_company_name(name: str) -> str:
    """公司名标准化：去空格、括号、标点和常见法律后缀。"""
    if name is None:
        return ""
    s = str(name).strip()
    s = s.replace("（", "(").replace("）", ")")
    s = re.sub(r"\([^)]*\)", "", s)  # 删除括号内信息，如（成都）
    s = re.sub(r"[\s\u3000·•,，.。;；:：、\-—_（）()\[\]【】]", "", s)
    for suf in COMPANY_LEGAL_SUFFIXES:
        if s.endswith(suf):
            s = s[: -len(suf)]
    return s


def strip_location_prefix(s: str) -> str:
    for pre in COMMON_LOCATION_PREFIXES:
        if s.startswith(pre) and len(s) > len(pre) + 1:
            return s[len(pre):]
    return s


def core_company_candidates(name: str) -> List[str]:
    """
    为简称和全称生成多个候选核心词。
    例如：
    - 迈瑞医疗 -> 迈瑞医疗 / 迈瑞
    - 深圳迈瑞生物医疗电子 -> 迈瑞生物医疗电子 / 迈瑞生物 / 迈瑞
    """
    norm = normalize_company_name(name)
    norm2 = strip_location_prefix(norm)

    cands = []
    for x in [norm, norm2]:
        if x and x not in cands:
            cands.append(x)

    # 对末尾行业词逐步剥离，形成核心品牌词
    for base in list(cands):
        cur = base
        changed = True
        while changed:
            changed = False
            for suf in COMPANY_INDUSTRY_SUFFIXES:
                if cur.endswith(suf) and len(cur) > len(suf) + 1:
                    cur = cur[: -len(suf)]
                    if cur and cur not in cands:
                        cands.append(cur)
                    changed = True
                    break

    # 对全称中出现的行业词前缀也做一次切割：深圳迈瑞生物医疗电子 -> 深圳迈瑞 / 迈瑞
    for base in list(cands):
        for suf in COMPANY_INDUSTRY_SUFFIXES:
            pos = base.find(suf)
            if pos >= 2:
                left = base[:pos]
                left = strip_location_prefix(left)
                if len(left) >= 2 and left not in cands:
                    cands.append(left)

    # 只保留长度>=2的候选，短到1个字容易误匹配
    return [x for x in cands if len(x) >= 2]


def lcs_len(a: str, b: str) -> int:
    """最长公共子序列长度，用于判断“迈瑞医疗”这类非连续匹配。"""
    if not a or not b:
        return 0
    # a 通常较短；为了节省内存，只保留一行 DP
    prev = [0] * (len(b) + 1)
    for ca in a:
        cur = [0]
        for j, cb in enumerate(b, start=1):
            if ca == cb:
                cur.append(prev[j - 1] + 1)
            else:
                cur.append(max(prev[j], cur[-1]))
        prev = cur
    return prev[-1]


def company_match_score(query_name: str, candidate_name: str) -> Tuple[float, str]:
    """
    返回公司名匹配分数和原因。
    分数设计偏保守：只有核心品牌词明确命中或字符覆盖率很高才通过。
    """
    q = normalize_company_name(query_name)
    c = normalize_company_name(candidate_name)
    if not q or not c:
        return 0.0, "empty"

    if q == c:
        return 100.0, "exact"

    if q in c or c in q:
        # 简称直接包含在全称中，如“华健未来” in “华健未来成都科技”
        return 96.0, "contains"

    q_cores = core_company_candidates(q)
    c_cores = core_company_candidates(c)

    # 核心词包含匹配：迈瑞医疗 -> 迈瑞；铂科电子 -> 铂科
    for qc in q_cores:
        if len(qc) >= 2 and qc in c:
            # 2字核心词给 92，3字以上给更高；防止过宽松
            return (94.0 if len(qc) >= 3 else 92.0), f"query_core_in_candidate:{qc}"

    for qc in q_cores:
        for cc in c_cores:
            if qc == cc and len(qc) >= 2:
                return (95.0 if len(qc) >= 3 else 92.0), f"same_core:{qc}"
            if len(qc) >= 3 and (qc in cc or cc in qc):
                return 90.0, f"core_contains:{qc}/{cc}"

    # 字符覆盖率：query 中多少非重复字符出现在 candidate 中
    q_chars = [ch for ch in q if "\u4e00" <= ch <= "\u9fff"]
    q_set = set(q_chars)
    c_set = set([ch for ch in c if "\u4e00" <= ch <= "\u9fff"])
    if q_set:
        coverage = len(q_set & c_set) / len(q_set)
    else:
        coverage = 0.0

    lcs = lcs_len(q, c)
    lcs_ratio = lcs / max(1, len(q))

    # 迈瑞医疗 vs 深圳迈瑞生物医疗电子：coverage=1, lcs_ratio=1，可以匹配
    if coverage >= 0.95 and lcs_ratio >= 0.80 and len(q) >= 4:
        return 88.0, f"char_lcs:coverage={coverage:.2f},lcs={lcs_ratio:.2f}"

    # rapidfuzz 兜底，但阈值较高，避免误匹配
    wr = fuzz.WRatio(q, c)
    pr = fuzz.partial_ratio(q, c)
    score = max(wr, pr)
    if score >= 88:
        return float(score), f"fuzzy:{score:.0f}"

    return float(max(score, coverage * 80, lcs_ratio * 80)), f"low:fuzzy={score:.0f},coverage={coverage:.2f},lcs={lcs_ratio:.2f}"


def find_company_row(ws, company: str, name_col: int = 2, min_score: float = 88.0) -> Optional[int]:
    """兼容旧调用：只返回行号。"""
    row, _, _, _ = find_company_row_detail(ws, company, name_col=name_col, min_score=min_score)
    return row


def find_company_row_detail(ws, company: str, name_col: int = 2, min_score: float = 88.0) -> Tuple[Optional[int], Optional[str], float, str]:
    """
    多轮公司名匹配。
    返回：(行号, 匹配到的主表公司名, 分数, 匹配原因)
    如果低于阈值或结果不唯一，返回 None。
    """
    candidates = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, name_col).value
        if v:
            candidates.append((r, str(v).strip()))

    if not candidates:
        return None, None, 0.0, "no_candidates"

    scored = []
    for r, name in candidates:
        score, reason = company_match_score(company, name)
        scored.append((score, r, name, reason))

    scored.sort(reverse=True, key=lambda x: x[0])
    best_score, best_row, best_name, best_reason = scored[0]
    second_score = scored[1][0] if len(scored) > 1 else 0.0

    if best_score < min_score:
        return None, best_name, best_score, best_reason

    # 如果前两名非常接近，说明可能歧义，宁可跳过，不要写错公司
    if second_score >= min_score and best_score - second_score < 3:
        return None, best_name, best_score, f"ambiguous:{best_reason}; second={scored[1][2]}({second_score:.1f})"

    return best_row, best_name, best_score, best_reason

def update_master_with_issues(master_path: str, issues: List[Issue], supp_date: str, out_path: str, supp_week_start: str = ""):
    supp_date = format_date_slash(supp_date)
    supp_week_start = format_date_slash(supp_week_start) if supp_week_start else ""
    """
    将补充材料问题写回主表。

    重要逻辑：
    1. docx 里出现的公司，必须先匹配到主表 Sheet1 / 备案情况总表中的公司，才写入。
    2. 匹配到的公司：I列状态改为“补充材料”，K列写补充材料公告日期，O列写问题数量，P列起写问题原文。
    3. 匹配不到的公司：整家公司跳过，不写 Sheet1，不追加“本次新增”，也不追加“问题类型统计”。
       例如“再惠网络”如果主表之前完全没有披露，就不会进入统计。
    4. 生成“公司匹配日志”sheet，方便人工检查哪些公司匹配成功/失败。
    """
    wb = safe_load_workbook(master_path)
    ws = wb["备案情况总表"] if "备案情况总表" in wb.sheetnames else wb.active
    history_ws = wb["问题类型统计"] if "问题类型统计" in wb.sheetnames else None
    add_ws = wb["本次新增"] if "本次新增" in wb.sheetnames else wb.create_sheet("本次新增")

    # 匹配日志：每次运行都重建，避免旧日志混淆
    if "公司匹配日志" in wb.sheetnames:
        del wb["公司匹配日志"]
    log_ws = wb.create_sheet("公司匹配日志")
    log_headers = ["公告公司名", "主表匹配公司名", "匹配分数", "匹配原因", "是否写入", "问题数量", "说明"]
    for c, h in enumerate(log_headers, start=1):
        log_ws.cell(1, c).value = h

    by_company: Dict[str, List[Issue]] = {}
    for it in issues:
        by_company.setdefault(it.company, []).append(it)

    matched_issues: List[Issue] = []
    log_row = 2

    # 备案情况总表：I为备案状态，K为补充材料公告日期，O为问题数量，P开始问题1
    for company, company_issues in by_company.items():
        r, matched_name, score, reason = find_company_row_detail(ws, company, name_col=2, min_score=88.0)

        if not r:
            log_ws.cell(log_row, 1).value = company
            log_ws.cell(log_row, 2).value = matched_name or ""
            log_ws.cell(log_row, 3).value = round(score, 1)
            log_ws.cell(log_row, 4).value = reason
            log_ws.cell(log_row, 5).value = "否"
            log_ws.cell(log_row, 6).value = len(company_issues)
            log_ws.cell(log_row, 7).value = "主表未匹配到该公司，已跳过；不会写入Sheet1/本次新增/问题类型统计"
            log_row += 1
            print(f"[SKIP] 主表未匹配到公司：{company}；best={matched_name} score={score:.1f} reason={reason}")
            continue

        # 匹配成功：写入主表
        ws.cell(r, 9).value = "补充材料"          # I: 备案状态 / 补充材料状态
        ws.cell(r, 10).value = supp_week_start     # J: 备案补充材料公告当周（文件名起始日期）
        ws.cell(r, 11).value = supp_date           # K: 备案补充材料公告日期（文件名结束日期）
        ws.cell(r, 15).value = len(company_issues)  # O: 问题数量

        # P列起：问题1、问题2……写“原问题”，不是标准化概括
        for i, it in enumerate(company_issues[:9], start=0):
            ws.cell(r, 16 + i).value = it.original

        # 如果本次问题少于9个，清空后面旧问题，避免上次残留
        for j in range(len(company_issues), 9):
            ws.cell(r, 16 + j).value = None

        matched_issues.extend(company_issues)

        log_ws.cell(log_row, 1).value = company
        log_ws.cell(log_row, 2).value = matched_name
        log_ws.cell(log_row, 3).value = round(score, 1)
        log_ws.cell(log_row, 4).value = reason
        log_ws.cell(log_row, 5).value = "是"
        log_ws.cell(log_row, 6).value = len(company_issues)
        log_ws.cell(log_row, 7).value = f"已写入主表第{r}行"
        log_row += 1

    # 本次新增：只追加“匹配成功公司”的问题
    start = add_ws.max_row + 1
    for it in matched_issues:
        add_ws.cell(start, 1).value = it.company
        add_ws.cell(start, 2).value = it.original
        add_ws.cell(start, 3).value = it.summary
        add_ws.cell(start, 4).value = it.category
        start += 1

    # 问题类型统计：只追加“匹配成功公司”的标准问题 + 大类，供下次复用
    if history_ws:
        hrow = history_ws.max_row + 1
        for it in matched_issues:
            history_ws.cell(hrow, 1).value = it.summary
            history_ws.cell(hrow, 2).value = it.category
            history_ws.cell(hrow, 3).value = 1
            hrow += 1

    wb.save(out_path)
    return out_path


def update_completed_filings(master_path: str, completed_items: List[Tuple[str, str]], out_path: str):
    """completed_items: [(公司名, 完成备案日期)]"""
    wb = safe_load_workbook(master_path)
    ws = wb["备案情况总表"] if "备案情况总表" in wb.sheetnames else wb.active
    for company, completed_date in completed_items:
        r = find_company_row(ws, company, name_col=2)
        if r:
            ws.cell(r, 9).value = "已完成备案"   # I: 备案状态
            ws.cell(r, 12).value = format_date_slash(completed_date) # L: 已完成备案时间
        else:
            print(f"[WARN] 未找到完成备案公司：{company}")
    wb.save(out_path)
    return out_path


def sync_filing_xlsx(master_path: str, new_filing_xlsx: str, out_path: str):
    """把官网最新备案情况表中的新增企业或更新信息同步到主表。"""
    master_wb = safe_load_workbook(master_path)
    master_ws = master_wb["备案情况总表"] if "备案情况总表" in master_wb.sheetnames else master_wb.active
    new_wb = safe_load_workbook(new_filing_xlsx, data_only=False)
    new_ws = new_wb.active

    # 新表列 A-J，对应主表 A-I 基本信息。主表 C-I 与新表 C-I 基本一致。
    master_names = {str(master_ws.cell(r, 2).value).strip(): r for r in range(5, master_ws.max_row+1) if master_ws.cell(r, 2).value}
    append_row = master_ws.max_row + 1
    changed = []

    for r in range(5, new_ws.max_row+1):
        name = new_ws.cell(r, 2).value
        if not name:
            continue
        name_s = str(name).strip()
        # 精确或模糊匹配
        if name_s in master_names:
            mr = master_names[name_s]
        else:
            best = process.extractOne(name_s, list(master_names.keys()), scorer=fuzz.token_set_ratio) if master_names else None
            mr = master_names[best[0]] if best and best[1] >= 92 else None
        if mr:
            # 更新 C-I 基础字段；不覆盖后面的补充材料和问题列
            for c in range(3, 10):
                old = master_ws.cell(mr, c).value
                new = new_ws.cell(r, c).value
                if new is not None and old != new:
                    master_ws.cell(mr, c).value = new
                    changed.append(("update", name_s, c))
        else:
            # 追加新企业，复制 A-I
            for c in range(1, 10):
                master_ws.cell(append_row, c).value = new_ws.cell(r, c).value
            changed.append(("append", name_s, None))
            append_row += 1

    master_wb.save(out_path)
    return changed


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--master", required=True, help="你的总表 xlsx")
    parser.add_argument("--supp-doc", help="补充材料要求 docx")
    parser.add_argument("--filing-xlsx", help="官网备案情况表 xlsx")
    parser.add_argument("--supp-date", default=datetime.today().strftime("%Y/%-m/%-d") if os.name != "nt" else datetime.today().strftime("%Y/%#m/%#d"))
    parser.add_argument("--supp-week-start", default="")
    parser.add_argument("--out", default="output.xlsx")
    args = parser.parse_args()

    current = args.master
    if args.filing_xlsx:
        tmp = args.out.replace(".xlsx", "_synced.xlsx")
        changed = sync_filing_xlsx(current, args.filing_xlsx, tmp)
        print(f"同步备案情况表完成：{len(changed)} 处新增/更新")
        current = tmp

    if args.supp_doc:
        issues = parse_supplement_doc(args.supp_doc, current)
        print(f"解析补充材料完成：{len(issues)} 个大问题")
        for it in issues[:10]:
            print(it.company, "|", it.summary, "|", it.category)
        update_master_with_issues(current, issues, args.supp_date, args.out, supp_week_start=args.supp_week_start)
        print(f"已输出：{args.out}")

if __name__ == "__main__":
    main()
